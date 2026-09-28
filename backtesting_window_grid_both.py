import numpy as np
import pandas as pd
from pathlib import Path
from zoneinfo import ZoneInfo
import re
import time
from io import StringIO
from openpyxl.formatting.rule import ColorScaleRule
import openpyxl.utils

DATA_DIR = Path(".")
OUTPUT_DIR = DATA_DIR / "results"

TP_VALUES = [x / 1000 for x in range(1, 8)]
SL_VALUES = [x / 1000 for x in range(1, 8)]

SAME_CANDLE_RULE = "sl_first"
ESTIMATED_ROLLOVER_FEE_PERCENT_PER_DAY = 0.01
ROLLOVER_HOUR_UTC = 0

ASSET_TIMEZONES = {
    "USD": "America/New_York",
    "GBP": "Europe/London",
    "EUR": "Europe/Berlin",
    "JPY": "Asia/Tokyo",
    "AUD": "Australia/Sydney",
    "NZD": "Pacific/Auckland",
    "CHF": "Europe/Zurich",
    "CAD": "America/Toronto",
}

pair_sets = {
    "USD": ["EURUSD", "GBPUSD", "USDJPY", "AUDUSD", "NZDUSD", "USDCAD", "USDCHF"],
    "EUR": ["EURUSD", "EURJPY", "EURGBP", "EURAUD", "EURNZD", "EURCHF", "EURCAD"],
    "GBP": ["GBPUSD", "GBPJPY", "EURGBP", "GBPAUD", "GBPNZD", "GBPCHF", "GBPCAD"],
    "JPY": ["USDJPY", "EURJPY", "GBPJPY", "AUDJPY", "NZDJPY", "CADJPY", "CHFJPY"],
    "AUD": ["AUDUSD", "AUDJPY", "EURAUD", "GBPAUD", "AUDNZD", "AUDCHF", "AUDCAD"],
    "NZD": ["NZDUSD", "NZDJPY", "EURNZD", "GBPNZD", "AUDNZD", "NZDCHF", "NZDCAD"],
    "CHF": ["USDCHF", "EURCHF", "GBPCHF", "AUDCHF", "NZDCHF", "CADCHF", "CHFJPY"],
    "CAD": ["USDCAD", "EURCAD", "GBPCAD", "AUDCAD", "NZDCAD", "CADCHF", "CADJPY"],
}


def choose_currency_group():
    print("\nAvailable groups:", ", ".join(pair_sets.keys()))
    choice = input("Type currency group (example: GBP): ").upper().strip()
    if choice in pair_sets:
        return choice
    print("Invalid currency group.")
    return None


def choose_entry_time():
    return input("Enter trade entry time in local event timezone HH:MM: ").strip()


def choose_inverted_event():
    print("\nIs higher Actual worse for the currency? (y/n)")
    choice = input("Choice: ").lower().strip()
    return choice == "y"


def choose_grid_settings():
    print("\n--- Window Grid Search Settings ---")
    min_train = int(input("Min Train Window Size (default 3): ").strip() or 3)
    max_train = int(input("Max Train Window Size (default 13): ").strip() or 13)
    
    min_fwd = int(input("Min Forward Test Window Size (default 3): ").strip() or 3)
    max_fwd = int(input("Max Forward Test Window Size (default 13): ").strip() or 13)

    cutoff_str = input("Top Lowest Delta Percentage Cutoff (e.g., 15 or 30) (default 50): ").strip()
    top_pct = float(cutoff_str) / 100.0 if cutoff_str.replace(".", "", 1).isdigit() else 0.50

    min_wr_str = input("Minimum Acceptable Win Rate % (e.g., 50) (default 50): ").strip()
    min_win_rate = float(min_wr_str) if min_wr_str.replace(".", "", 1).isdigit() else 50.0

    return min_train, max_train, min_fwd, max_fwd, top_pct, min_win_rate


def load_csv(pair):
    file_path = DATA_DIR / f"{pair}60.csv"
    if not file_path.exists():
        print(f"Missing CSV for {pair}: {file_path}")
        return None
    df = pd.read_csv(file_path, sep="\t", header=None, names=["Datetime", "Open", "High", "Low", "Close", "Volume"], engine="c")
    required_cols = ["Datetime", "Open", "High", "Low", "Close"]
    missing_cols = [col for col in required_cols if col not in df.columns]
    if missing_cols:
        print(f"{pair} missing columns: {missing_cols}")
        return None
    df["Datetime"] = pd.to_datetime(df["Datetime"], utc=True, errors="coerce").dt.tz_localize(None)
    for col in ["Open", "High", "Low", "Close"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df = (
        df.dropna(subset=required_cols)
        .sort_values("Datetime")
        .drop_duplicates("Datetime")
        .set_index("Datetime")
    )
    return df


def get_direction(mode, pair, surprise, event_currency, candle=None, sub_mode=None):
    if mode == "surprise":
        if surprise == "neutral":
            return None
        if pair.startswith(event_currency):
            return "long" if surprise == "positive" else "short"
        if pair.endswith(event_currency):
            return "short" if surprise == "positive" else "long"
    elif mode == "candle_colour":
        if candle is None:
            return None
        if candle["Close"] > candle["Open"]:
            candle_colour = "green"
        elif candle["Close"] < candle["Open"]:
            candle_colour = "red"
        else:
            return None
        if sub_mode == "trend":
            return "long" if candle_colour == "green" else "short"
        if sub_mode == "fade":
            return "short" if candle_colour == "green" else "long"
    return None


def paste_news_data_once(event_currency, entry_time_str, inverted_event):
    print("\n--- Paste News Data (used for BOTH modes) ---")
    print("Paste your tab-separated news data below (with header row).")
    print("Must include: History, Actual, Forecast columns.")
    print("When done pasting, enter an empty line:\n")
    lines = []
    while True:
        line = input()
        if line.strip() == "" and lines:
            break
        lines.append(line)
    if not lines:
        print("No data pasted.")
        return None
    raw = "\n".join(lines)
    df = pd.read_csv(StringIO(raw), sep="\t")
    if "History" not in df.columns:
        print("Missing 'History' column.")
        return None
    if "Actual" not in df.columns or "Forecast" not in df.columns:
        print("Missing 'Actual' or 'Forecast' columns (required for surprise mode).")
        return None
    try:
        hour, minute = map(int, entry_time_str.split(":"))
    except ValueError:
        print("Invalid time format. Use HH:MM.")
        return None

    def parse_history_date(date_str):
        date_str = str(date_str).strip()
        date_str = re.sub(r'(\w+\s+\d+)-(\d+),\s*(\d{4})', r'\1, \3', date_str)
        return pd.to_datetime(date_str, format="%b %d, %Y", errors="coerce")

    df["History"] = df["History"].astype(str).apply(parse_history_date)

    def parse_value(val):
        val = str(val).replace("%", "").replace(",", "").strip().lower()
        if val.endswith("k") or val.endswith("m"):
            return val[:-1]
        return val

    for col in ["Actual", "Forecast"]:
        df[col] = pd.to_numeric(df[col].astype(str).apply(parse_value), errors="coerce")
    df = df.dropna(subset=["History", "Actual", "Forecast"]).copy()

    local_tz = ZoneInfo(ASSET_TIMEZONES[event_currency])
    utc_tz = ZoneInfo("UTC")

    df["UTC_Time"] = [
        pd.Timestamp(
            year=d.year, month=d.month, day=d.day,
            hour=hour, minute=minute, tz=local_tz,
        ).tz_convert(utc_tz).tz_localize(None)
        for d in df["History"]
    ]

    adjusted_surprise = (df["Actual"] - df["Forecast"]) * (-1 if inverted_event else 1)
    df["Surprise_Type"] = adjusted_surprise.apply(
        lambda value: "positive" if value > 0 else "negative" if value < 0 else "neutral"
    )

    return df.sort_values("UTC_Time").reset_index(drop=True)


def prepare_candle_news(news_df):
    candle_df = news_df[["History", "UTC_Time"]].copy()
    candle_df["Surprise_Type"] = "neutral"
    return candle_df


def calculate_pnl(entry_price, exit_price, direction):
    if direction == "long":
        return (exit_price - entry_price) / entry_price
    return (entry_price - exit_price) / entry_price


def calculate_rollover_fee_percent(entry_time, exit_time):
    entry_day = (entry_time - pd.Timedelta(hours=ROLLOVER_HOUR_UTC)).date()
    exit_day = (exit_time - pd.Timedelta(hours=ROLLOVER_HOUR_UTC)).date()
    rollover_count = max(0, (exit_day - entry_day).days)
    return rollover_count * ESTIMATED_ROLLOVER_FEE_PERCENT_PER_DAY


def build_trade_result(entry_price, exit_price, direction, entry_time, exit_time):
    hold_hours = (exit_time - entry_time).total_seconds() / 3600
    gross_pnl_pct = calculate_pnl(entry_price, exit_price, direction)
    rollover_fee_percent = calculate_rollover_fee_percent(entry_time, exit_time)
    rollover_fee_pct = rollover_fee_percent / 100
    return {
        "pnl_pct": gross_pnl_pct - rollover_fee_pct,
        "gross_pnl_pct": gross_pnl_pct,
        "rollover_fee_percent": rollover_fee_percent,
        "hold_hours": hold_hours,
        "entry_price": entry_price,
        "exit_price": exit_price,
    }


def get_end_of_week_cutoff(entry_time):
    days_ahead = 4 - entry_time.weekday()
    if days_ahead < 0 or (days_ahead == 0 and entry_time.hour >= 18):
        days_ahead += 7
    target = entry_time + pd.Timedelta(days=days_ahead)
    return target.replace(hour=18, minute=0, second=0, microsecond=0)


def simulate_trade(df, entry_time, direction, tp_pct, sl_pct):
    if entry_time not in df.index:
        return None
    entry_price = float(df.loc[entry_time, "Open"])
    if direction == "long":
        tp_price = entry_price * (1 + tp_pct)
        sl_price = entry_price * (1 - sl_pct)
    else:
        tp_price = entry_price * (1 - tp_pct)
        sl_price = entry_price * (1 + sl_pct)

    cutoff = get_end_of_week_cutoff(entry_time)
    available = df.index[df.index <= cutoff]
    if available.empty:
        return None
    window_end = available[-1]
    window = df.loc[entry_time:window_end]

    if window.empty:
        return None

    for current_time, row in window.iterrows():
        if direction == "long":
            tp_hit = row["High"] >= tp_price
            sl_hit = row["Low"] <= sl_price
        else:
            tp_hit = row["Low"] <= tp_price
            sl_hit = row["High"] >= sl_price

        if tp_hit and sl_hit:
            if SAME_CANDLE_RULE == "sl_first":
                res = build_trade_result(entry_price, sl_price, direction, entry_time, current_time)
                res["exit_reason"] = "SL (Same Candle)"
                return res
            res = build_trade_result(entry_price, tp_price, direction, entry_time, current_time)
            res["exit_reason"] = "TP (Same Candle)"
            return res
        if tp_hit:
            res = build_trade_result(entry_price, tp_price, direction, entry_time, current_time)
            res["exit_reason"] = "TP"
            return res
        if sl_hit:
            res = build_trade_result(entry_price, sl_price, direction, entry_time, current_time)
            res["exit_reason"] = "SL"
            return res

    manual_exit_price = float(window.iloc[-1]["Close"])
    res = build_trade_result(entry_price, manual_exit_price, direction, entry_time, window.index[-1])
    res["exit_reason"] = "Friday Cutoff"
    return res


def calculate_max_drawdown(pnl_series):
    if len(pnl_series) == 0:
        return 0.0
    cum_pnl = np.cumsum(pnl_series)
    running_max = np.maximum.accumulate(cum_pnl)
    drawdowns = running_max - cum_pnl
    return float(np.max(drawdowns))


def run_backtest(mode, news_df, pair_dfs, group, min_train, max_train, min_fwd, max_fwd, top_pct_cutoff, min_acceptable_win_rate):
    print(f"\n{'='*60}")
    print(f"  RUNNING BACKTEST: {mode.upper()} MODE")
    print(f"{'='*60}")

    if mode == "surprise":
        work_df = news_df[news_df["Surprise_Type"] != "neutral"].reset_index(drop=True)
        print(f"Filtered out neutral surprise events. Usable events: {len(work_df)}")
    else:
        valid_mask = []
        for _, row in news_df.iterrows():
            entry_time_utc = row["UTC_Time"]
            prev_time_utc = entry_time_utc - pd.Timedelta(hours=1)
            is_valid = any(
                entry_time_utc in pdf.index and prev_time_utc in pdf.index
                for pdf in pair_dfs.values()
            )
            valid_mask.append(is_valid)
        work_df = news_df[valid_mask].reset_index(drop=True)
        print(f"Filtered out events with no valid candle. Usable events: {len(work_df)}")

    total_events = len(work_df)
    if total_events == 0:
        print("No usable events after filtering.")
        return None, None, None

    print("\nPre-computing trade simulation lookup table...")
    
    sub_modes = ["trend", "fade"] if mode == "candle_colour" else ["N/A"]
    
    trade_cache = {}
    news_rows = work_df.to_dict("records")
    
    for event_idx, nr in enumerate(news_rows):
        entry_time_utc = nr["UTC_Time"]
        surprise_type = nr["Surprise_Type"]

        for pair, pdf in pair_dfs.items():
            if entry_time_utc not in pdf.index:
                continue

            if mode == "candle_colour":
                prev_time_utc = entry_time_utc - pd.Timedelta(hours=1)
                if prev_time_utc not in pdf.index:
                    continue
                candle = pdf.loc[prev_time_utc]
            else:
                candle = None

            for sub_mode in sub_modes:
                direction = get_direction(
                    mode=mode, pair=pair, surprise=surprise_type,
                    event_currency=group, candle=candle,
                    sub_mode=sub_mode if sub_mode != "N/A" else None,
                )
                if direction is None:
                    continue

                for tp in TP_VALUES:
                    for sl in SL_VALUES:
                        res = simulate_trade(pdf, entry_time_utc, direction, tp, sl)
                        if res is not None:
                            trade_cache[(event_idx, pair, sub_mode, tp, sl)] = {
                                "direction": direction,
                                "pnl_pct": res["pnl_pct"],
                                "hold_hours": res["hold_hours"],
                                "rollover_fee_percent": res["rollover_fee_percent"],
                                "entry_price": res["entry_price"],
                                "exit_price": res["exit_price"],
                                "exit_reason": res["exit_reason"],
                            }

    print(f"Pre-computation complete! Cached combinations: {len(trade_cache):,}")

    def evaluate_strategy_on_slice(pair, sub_mode, tp, sl, start_idx, end_idx):
        trades_pnl = []
        trades_hold = []
        trades_rollover = []

        for idx in range(start_idx, end_idx):
            key = (idx, pair, sub_mode, tp, sl)
            if key in trade_cache:
                t = trade_cache[key]
                trades_pnl.append(t["pnl_pct"])
                trades_hold.append(t["hold_hours"])
                trades_rollover.append(t["rollover_fee_percent"])

        count = len(trades_pnl)
        if count == 0:
            return {
                "pair": pair, "sub_mode": sub_mode,
                "tp_percent": tp * 100, "sl_percent": sl * 100,
                "risk_to_reward_ratio": tp / sl,
                "trade_count": 0, "win_rate_percent": 0.0,
                "total_pnl_percent": 0.0, "average_pnl_percent": 0.0,
                "average_hold_hours": 0.0,
                "total_rollover_fee_percent": 0.0,
                "average_rollover_fee_percent": 0.0,
            }

        pnl_arr = np.array(trades_pnl)
        hold_arr = np.array(trades_hold)
        roll_arr = np.array(trades_rollover)

        return {
            "pair": pair, "sub_mode": sub_mode,
            "tp_percent": tp * 100, "sl_percent": sl * 100,
            "risk_to_reward_ratio": tp / sl,
            "trade_count": count,
            "win_rate_percent": float((pnl_arr > 0).sum()) / count * 100,
            "total_pnl_percent": float(pnl_arr.sum()) * 100,
            "average_pnl_percent": float(pnl_arr.mean()) * 100,
            "average_hold_hours": float(hold_arr.mean()),
            "total_rollover_fee_percent": float(roll_arr.sum()),
            "average_rollover_fee_percent": float(roll_arr.mean()),
        }

    candidate_strategies = []
    for pair in pair_dfs.keys():
        for sub_mode in sub_modes:
            for tp in TP_VALUES:
                for sl in SL_VALUES:
                    candidate_strategies.append((pair, sub_mode, tp, sl))

    grid_results = []
    best_overall_log = []
    best_overall_pnl = -99999.0
    best_overall_window = None

    print("\nStarting Window Grid Search...")

    for train_size in range(min_train, max_train + 1):
        for fwd_size in range(min_fwd, max_fwd + 1):
            total_window = train_size + fwd_size
            min_req = total_window + 1
            
            if total_events < min_req:
                continue
                
            num_iterations = total_events - total_window
            current_trade_logs = []
            skipped_count = 0

            for i in range(num_iterations):
                train_start, train_end = i, i + train_size
                fwd_start, fwd_end = i + train_size, i + total_window
                target_idx = i + total_window

                target_event = work_df.iloc[target_idx]
                target_time_utc = target_event["UTC_Time"]
                history_date = target_event["History"]

                strategy_rows = []

                for pair, sub_mode, tp, sl in candidate_strategies:
                    train_metrics = evaluate_strategy_on_slice(pair, sub_mode, tp, sl, train_start, train_end)
                    fwd_metrics = evaluate_strategy_on_slice(pair, sub_mode, tp, sl, fwd_start, fwd_end)

                    if train_metrics["trade_count"] > 0 or fwd_metrics["trade_count"] > 0:
                        strategy_rows.append({
                            "pair": pair, "sub_mode": sub_mode, "tp": tp, "sl": sl,
                            "risk_to_reward_ratio": tp / sl,
                            "train_pnl": train_metrics["total_pnl_percent"],
                            "fwd_pnl": fwd_metrics["total_pnl_percent"],
                            "fwd_avg_pnl": fwd_metrics["average_pnl_percent"],
                            "avg_pnl_delta": abs(train_metrics["average_pnl_percent"] - fwd_metrics["average_pnl_percent"]),
                            "train_wr": train_metrics["win_rate_percent"],
                            "fwd_wr": fwd_metrics["win_rate_percent"]
                        })

                if not strategy_rows:
                    skipped_count += 1
                    current_trade_logs.append({
                        "Iteration": i + 1, "Date": history_date.strftime("%Y-%m-%d"),
                        "Pair": "None", "PnL_Percent": 0.0, "Is_Win": 0,
                        "Reason": "No candidate strategies"
                    })
                    continue

                df_strats = pd.DataFrame(strategy_rows)
                cutoff_count = max(1, int(np.ceil(len(df_strats) * top_pct_cutoff)))
                df_filtered = df_strats.sort_values("avg_pnl_delta").iloc[:cutoff_count]
                df_filtered = df_filtered.sort_values(["fwd_avg_pnl", "risk_to_reward_ratio"], ascending=[False, False])

                selected_strategy = None
                for _, row in df_filtered.iterrows():
                    if (row["train_wr"] >= min_acceptable_win_rate and
                        row["fwd_wr"] >= min_acceptable_win_rate and
                        row["train_pnl"] > 0 and row["fwd_pnl"] > 0):
                        selected_strategy = row
                        break

                if selected_strategy is None:
                    alt_sorted = df_strats.sort_values(["fwd_avg_pnl", "risk_to_reward_ratio"], ascending=[False, False])
                    for _, row in alt_sorted.iterrows():
                        if (row["train_wr"] >= min_acceptable_win_rate and
                            row["fwd_wr"] >= min_acceptable_win_rate and
                            row["train_pnl"] > 0 and row["fwd_pnl"] > 0):
                            selected_strategy = row
                            break

                if selected_strategy is not None:
                    trade_key = (target_idx, selected_strategy["pair"], selected_strategy["sub_mode"], selected_strategy["tp"], selected_strategy["sl"])
                    if trade_key in trade_cache:
                        trade = trade_cache[trade_key]
                        pnl_val = trade["pnl_pct"] * 100
                        current_trade_logs.append({
                            "Iteration": i + 1, "Date": history_date.strftime("%Y-%m-%d"),
                            "Pair": selected_strategy["pair"], "PnL_Percent": pnl_val,
                            "Is_Win": 1 if pnl_val > 0 else 0, "Reason": trade["exit_reason"]
                        })
                    else:
                        current_trade_logs.append({
                            "Iteration": i + 1, "Date": history_date.strftime("%Y-%m-%d"),
                            "Pair": "None", "PnL_Percent": 0.0, "Is_Win": 0,
                            "Reason": "Strategy outcome not cached"
                        })
                else:
                    skipped_count += 1
                    current_trade_logs.append({
                        "Iteration": i + 1, "Date": history_date.strftime("%Y-%m-%d"),
                        "Pair": "None", "PnL_Percent": 0.0, "Is_Win": 0,
                        "Reason": "No strategy with positive train & fwd PnL"
                    })

            trades_executed = [tl["PnL_Percent"] for tl in current_trade_logs if tl["Pair"] != "None"]
            total_trades = len(trades_executed)
            actual_total_pnl = sum(trades_executed)
            actual_win_rate = (sum([1 for pnl in trades_executed if pnl > 0]) / total_trades * 100) if total_trades > 0 else 0.0
            actual_ev = (actual_total_pnl / total_trades) if total_trades > 0 else 0.0
            max_dd = calculate_max_drawdown(trades_executed)

            grid_results.append({
                "Train_Size": train_size, "Forward_Size": fwd_size,
                "Train_Ratio_Percent": (train_size / total_window) * 100,
                "Total_Window": total_window, "Total_Trades": total_trades,
                "Skipped_Trades": skipped_count,
                "Actual_Win_Rate_Percent": actual_win_rate,
                "Actual_Total_PnL_Percent": actual_total_pnl,
                "Actual_EV_Percent": actual_ev, "Max_Drawdown_Percent": max_dd
            })

            if actual_total_pnl > best_overall_pnl:
                best_overall_pnl = actual_total_pnl
                best_overall_log = current_trade_logs
                best_overall_window = (train_size, fwd_size)

    if not grid_results:
        print("No valid grid results produced.")
        return None, None, None

    print(f"\n--- {mode.upper()} BACKTEST FINISHED ---")
    print(f"Best Window: {best_overall_window} (Train/Forward)")
    total_skipped = sum(r["Skipped_Trades"] for r in grid_results)
    print(f"Total skipped: {total_skipped}")

    return grid_results, best_overall_log, best_overall_window


def main():
    group = choose_currency_group()
    if not group:
        return

    entry_time = choose_entry_time()
    inverted_event = choose_inverted_event()

    min_train, max_train, min_fwd, max_fwd, top_pct_cutoff, min_acceptable_win_rate = choose_grid_settings()

    # --- PASTE NEWS DATA ONCE ---
    news_df = paste_news_data_once(group, entry_time, inverted_event)
    if news_df is None:
        return

    # --- LOAD PRICE DATA ONCE ---
    print("\nLoading price data for all pairs...")
    pair_dfs = {}
    for pair in pair_sets[group]:
        df = load_csv(pair)
        if df is not None:
            pair_dfs[pair] = df

    if not pair_dfs:
        print("No valid pair data found.")
        return

    # Filter events that have price data (do this ONCE for both modes)
    total_pasted = len(news_df)
    price_fit_mask = []
    for _, row in news_df.iterrows():
        entry_time_utc = row["UTC_Time"]
        has_price = any(entry_time_utc in pdf.index for pdf in pair_dfs.values())
        price_fit_mask.append(has_price)
    news_df = news_df[price_fit_mask].reset_index(drop=True)
    out_of_range = total_pasted - len(news_df)
    print(f"\nPasted data rows: {total_pasted}")
    print(f"Out of price data range: {out_of_range}")
    print(f"Events matching price data: {len(news_df)}")

    # Prepare candle mode news (only needs History + UTC_Time)
    candle_news_df = prepare_candle_news(news_df)

    # --- RUN SURPRISE MODE ---
    surprise_grid, surprise_log, surprise_window = run_backtest(
        "surprise", news_df, pair_dfs, group,
        min_train, max_train, min_fwd, max_fwd, top_pct_cutoff, min_acceptable_win_rate
    )

    # --- RUN CANDLE COLOUR MODE ---
    candle_grid, candle_log, candle_window = run_backtest(
        "candle_colour", candle_news_df, pair_dfs, group,
        min_train, max_train, min_fwd, max_fwd, top_pct_cutoff, min_acceptable_win_rate
    )

    if surprise_grid is None and candle_grid is None:
        print("Both modes produced no results.")
        return

    OUTPUT_DIR.mkdir(exist_ok=True)
    timestamp = time.strftime("%Y%m%d_%H%M%S")
    out_name = f"{group}_{timestamp}_window_grid_search_both.xlsx"
    out_path = OUTPUT_DIR / out_name

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # === SURPRISE MODE SHEETS ===
        if surprise_grid is not None:
            surprise_grid_df = pd.DataFrame(surprise_grid).sort_values("Actual_Total_PnL_Percent", ascending=False)
            surprise_log_df = pd.DataFrame(surprise_log)
            surprise_heatmap_df = pd.DataFrame(surprise_grid).pivot(index="Train_Size", columns="Forward_Size", values="Actual_Total_PnL_Percent")

            surprise_grid_df.to_excel(writer, sheet_name="Surprise_Grid_Summary", index=False)
            surprise_heatmap_df.to_excel(writer, sheet_name="Surprise_PnL_Heatmap")
            surprise_log_df.to_excel(writer, sheet_name="Surprise_Best_Trade_Log", index=False)
        else:
            pd.DataFrame().to_excel(writer, sheet_name="Surprise_Grid_Summary")
            pd.DataFrame(["No surprise mode results"]).to_excel(writer, sheet_name="Surprise_PnL_Heatmap")

        # === CANDLE COLOUR MODE SHEETS ===
        if candle_grid is not None:
            candle_grid_df = pd.DataFrame(candle_grid).sort_values("Actual_Total_PnL_Percent", ascending=False)
            candle_log_df = pd.DataFrame(candle_log)
            candle_heatmap_df = pd.DataFrame(candle_grid).pivot(index="Train_Size", columns="Forward_Size", values="Actual_Total_PnL_Percent")

            candle_grid_df.to_excel(writer, sheet_name="Candle_Grid_Summary", index=False)
            candle_heatmap_df.to_excel(writer, sheet_name="Candle_PnL_Heatmap")
            candle_log_df.to_excel(writer, sheet_name="Candle_Best_Trade_Log", index=False)
        else:
            pd.DataFrame().to_excel(writer, sheet_name="Candle_Grid_Summary")
            pd.DataFrame(["No candle colour mode results"]).to_excel(writer, sheet_name="Candle_PnL_Heatmap")

        # === APPLY CONDITIONAL FORMATTING TO BOTH HEATMAPS ===
        for sheet_name in ["Surprise_PnL_Heatmap", "Candle_PnL_Heatmap"]:
            if sheet_name in writer.sheets:
                ws = writer.sheets[sheet_name]
                max_col = ws.max_column
                max_row = ws.max_row
                col_letter = openpyxl.utils.get_column_letter(max_col)
                cell_range = f"B2:{col_letter}{max_row}"
                ws.conditional_formatting.add(
                    cell_range,
                    ColorScaleRule(
                        start_type='num', start_value=-10, start_color='F8696B',
                        mid_type='num', mid_value=0, mid_color='FFFFFF',
                        end_type='num', end_value=10, end_color='63BE7B'
                    )
                )

    print(f"\n{'='*60}")
    print(f"  ALL BACKTESTS COMPLETE")
    print(f"{'='*60}")
    print(f"File saved: {out_path}")
    print(f"\nSheets in file:")
    print(f"  1. Surprise_Grid_Summary")
    print(f"  2. Surprise_PnL_Heatmap")
    print(f"  3. Surprise_Best_Trade_Log")
    print(f"  4. Candle_Grid_Summary")
    print(f"  5. Candle_PnL_Heatmap")
    print(f"  6. Candle_Best_Trade_Log")

    if surprise_window:
        print(f"\nSurprise Mode - Best Window: {surprise_window} (Train/Forward)")
    if candle_window:
        print(f"Candle Colour Mode - Best Window: {candle_window} (Train/Forward)")


if __name__ == "__main__":
    main()