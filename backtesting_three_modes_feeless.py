"""
FEELESS copy of backtesting_three_modes.py - identical logic, but NO rollover/swap fee.

Runs THREE modes from one news paste, for one currency group:
  1. Candle Colour (trend + fade tested together)
  2. Surprise - higher Actual is BETTER for the currency
  3. Surprise - higher Actual is WORSE for the currency

Output: one Excel file
  - PnL_Heatmaps  : the three heatmaps side by side on ONE sheet
  - Summary_*     : window-grid summary per mode
  - Log_*         : best-window trade log per mode

Needs backtesting_window_grid.py in the same folder (reuses its settings,
pair lists, CSV loader, trade simulator). Edit TP/SL/pairs there.
"""
import re
import time
from io import StringIO
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import openpyxl.utils
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font

import backtesting_window_grid as bw


# ----------------------------------------------------------------------------
# Feeless override: rollover/swap fee always 0.
# backtesting_window_grid.build_trade_result() calls this function by name,
# so replacing it here removes the fee from every simulated trade.
# ----------------------------------------------------------------------------
def calculate_rollover_fee_percent(entry_time, exit_time):
    return 0.0


bw.calculate_rollover_fee_percent = calculate_rollover_fee_percent


MODES = [
    ("candle_colour", False, "Candle"),
    ("surprise", False, "Surprise_Better"),   # higher actual = better
    ("surprise", True, "Surprise_Worse"),     # higher actual = worse
]


# ----------------------------------------------------------------------------
# Input
# ----------------------------------------------------------------------------
def paste_news(group, entry_time_str):
    print("\n--- Paste News Data ---")
    print("Paste tab-separated news data (with header row). Empty line when done:\n")
    lines = []
    while True:
        line = input()
        if line.strip() == "" and lines:
            break
        lines.append(line)
    if not lines:
        print("No data pasted.")
        return None

    df = pd.read_csv(StringIO("\n".join(lines)), sep="\t")
    if "History" not in df.columns:
        print("Missing 'History' column.")
        return None
    try:
        hour, minute = map(int, entry_time_str.split(":"))
    except ValueError:
        print("Invalid time format. Use HH:MM.")
        return None

    def parse_date(s):
        s = str(s).strip()
        s = re.sub(r"(\w+\s+\d+)-(\d+),\s*(\d{4})", r"\1, \3", s)
        return pd.to_datetime(s, format="%b %d, %Y", errors="coerce")

    df["History"] = df["History"].astype(str).apply(parse_date)
    df = df.dropna(subset=["History"]).copy()

    has_surprise = "Actual" in df.columns and "Forecast" in df.columns
    if has_surprise:
        def parse_val(v):
            v = str(v).replace("%", "").replace(",", "").strip().lower()
            return v[:-1] if v.endswith(("k", "m")) else v
        for col in ["Actual", "Forecast"]:
            df[col] = pd.to_numeric(df[col].astype(str).apply(parse_val), errors="coerce")
        df["Raw_Diff"] = df["Actual"] - df["Forecast"]
    else:
        print("WARNING: no Actual/Forecast columns -> surprise modes will be skipped.")

    local_tz = ZoneInfo(bw.ASSET_TIMEZONES[group])
    df["UTC_Time"] = [
        pd.Timestamp(year=d.year, month=d.month, day=d.day, hour=hour,
                     minute=minute, tz=local_tz).tz_convert("UTC").tz_localize(None)
        for d in df["History"]
    ]
    return df.sort_values("UTC_Time").reset_index(drop=True), has_surprise


def prepare_events(news_df, mode, inverted, pair_arrays_map):
    """Filter/label events for one mode."""
    df = news_df.copy()

    if mode == "surprise":
        df = df.dropna(subset=["Actual", "Forecast"])
        adj = df["Raw_Diff"] * (-1 if inverted else 1)
        df["Surprise_Type"] = np.where(adj > 0, "positive", np.where(adj < 0, "negative", "neutral"))
        df = df[df["Surprise_Type"] != "neutral"]
        keep = [any(bw.has_candle(pa, t) for pa in pair_arrays_map.values()) for t in df["UTC_Time"]]
    else:
        df["Surprise_Type"] = "neutral"
        one_h = pd.Timedelta(hours=1)
        keep = [
            any(bw.has_candle(pa, t) and bw.has_candle(pa, t - one_h) for pa in pair_arrays_map.values())
            for t in df["UTC_Time"]
        ]
    return df[keep].reset_index(drop=True)


# ----------------------------------------------------------------------------
# Core grid search for ONE mode (same logic as backtesting_window_grid.py)
# ----------------------------------------------------------------------------
def run_grid(mode, news_df, group, pair_arrays_map, settings, sub_modes=None):
    min_train, max_train, min_fwd, max_fwd, top_pct, min_wr = settings
    total_events = len(news_df)
    if total_events == 0:
        return None, None, None

    if sub_modes is None:
        sub_modes = ["trend", "fade"] if mode == "candle_colour" else ["N/A"]
    elif mode != "candle_colour":
        sub_modes = ["N/A"]
    keys = [(p, s, tp, sl) for p in pair_arrays_map for s in sub_modes
            for tp in bw.TP_VALUES for sl in bw.SL_VALUES]
    cpos = {k: i for i, k in enumerate(keys)}
    n_c = len(keys)
    cand_rr = np.array([k[2] / k[3] for k in keys], dtype=float)

    metrics = np.zeros((3, n_c, total_events))
    lookup = {}
    rows = news_df.to_dict("records")

    for ei, nr in enumerate(rows):
        t = nr["UTC_Time"]
        for pair, pa in pair_arrays_map.items():
            window = bw.build_trade_window(pa, t)
            if window is None:
                continue
            candle = None
            if mode == "candle_colour":
                pp = bw.index_position(pa["index"], t - pd.Timedelta(hours=1))
                if pp < 0:
                    continue
                candle = {"Open": float(pa["open"][pp]), "Close": float(pa["close"][pp])}
            for sm in sub_modes:
                d = bw.get_direction(mode=mode, pair=pair, surprise=nr["Surprise_Type"],
                                     event_currency=group, candle=candle,
                                     sub_mode=None if sm == "N/A" else sm)
                if d is None:
                    continue
                for tp in bw.TP_VALUES:
                    for sl in bw.SL_VALUES:
                        res = bw.simulate_trade(window, d, tp, sl)
                        ci = cpos[(pair, sm, tp, sl)]
                        metrics[0, ci, ei] = res["pnl_pct"] * 100
                        metrics[1, ci, ei] = 1.0
                        if res["pnl_pct"] > 0:
                            metrics[2, ci, ei] = 1.0
                        lookup[(ei, ci)] = (res["pnl_pct"] * 100, res["exit_reason"])

    prefix = np.zeros((3, n_c, total_events + 1))
    np.cumsum(metrics, axis=2, out=prefix[:, :, 1:])
    del metrics
    pre_pnl, pre_n, pre_w = prefix
    zeros_c = np.zeros(n_c)

    grid_results, best_log, best_pnl, best_win = [], [], -99999.0, None

    for train_size in range(min_train, max_train + 1):
        for fwd_size in range(min_fwd, max_fwd + 1):
            total_window = train_size + fwd_size
            if total_events < total_window + 1:
                continue

            logs, skipped = [], 0

            def skip(i, date, reason):
                logs.append({"Iteration": i + 1, "Date": date, "Pair": "None",
                             "PnL_Percent": 0.0, "Is_Win": 0, "Reason": reason})

            for i in range(total_events - total_window):
                ts, te = i, i + train_size
                fs, fe = te, i + total_window
                tgt = i + total_window
                date = rows[tgt]["History"].strftime("%Y-%m-%d")

                n_tr = pre_n[:, te] - pre_n[:, ts]
                n_fw = pre_n[:, fe] - pre_n[:, fs]
                pool = np.flatnonzero((n_tr > 0) | (n_fw > 0))
                if pool.size == 0:
                    skipped += 1
                    skip(i, date, "No candidate strategies")
                    continue

                tr_pnl = pre_pnl[:, te] - pre_pnl[:, ts]
                fw_pnl = pre_pnl[:, fe] - pre_pnl[:, fs]
                tot_pnl = tr_pnl + fw_pnl  # train + forward total PnL (used for ranking)
                tr_avg = np.divide(tr_pnl, n_tr, out=zeros_c.copy(), where=n_tr > 0)
                fw_avg = np.divide(fw_pnl, n_fw, out=zeros_c.copy(), where=n_fw > 0)
                tr_wr = np.divide(pre_w[:, te] - pre_w[:, ts], n_tr, out=zeros_c.copy(), where=n_tr > 0) * 100
                fw_wr = np.divide(pre_w[:, fe] - pre_w[:, fs], n_fw, out=zeros_c.copy(), where=n_fw > 0) * 100

                delta = np.abs(tr_avg - fw_avg)
                cut = max(1, int(np.ceil(pool.size * top_pct)))
                by_delta = pool[np.argsort(delta[pool], kind="stable")][:cut]
                ordered = by_delta[np.lexsort((-cand_rr[by_delta], -tot_pnl[by_delta]))]

                ok = (tr_wr >= min_wr) & (fw_wr >= min_wr) & (tr_pnl > 0) & (fw_pnl > 0)

                picked = -1
                if ok[ordered].any():
                    picked = int(ordered[int(np.argmax(ok[ordered]))])
                else:
                    fb = pool[np.lexsort((-cand_rr[pool], -tot_pnl[pool]))]
                    if ok[fb].any():
                        picked = int(fb[int(np.argmax(ok[fb]))])

                if picked < 0:
                    skipped += 1
                    skip(i, date, "No strategy with positive train & fwd PnL")
                    continue

                cached = lookup.get((tgt, picked))
                if cached is None:
                    skipped += 1
                    skip(i, date, "Strategy outcome not cached")
                    continue

                pnl_val, reason = cached
                logs.append({"Iteration": i + 1, "Date": date, "Pair": keys[picked][0],
                             "PnL_Percent": pnl_val, "Is_Win": int(pnl_val > 0), "Reason": reason})

            tr = [l["PnL_Percent"] for l in logs if l["Pair"] != "None"]
            n = len(tr)
            total = sum(tr)
            grid_results.append({
                "Train_Size": train_size, "Forward_Size": fwd_size,
                "Train_Ratio_Percent": train_size / total_window * 100,
                "Total_Window": total_window, "Total_Trades": n, "Skipped_Trades": skipped,
                "Actual_Win_Rate_Percent": (sum(p > 0 for p in tr) / n * 100) if n else 0.0,
                "Actual_Total_PnL_Percent": total,
                "Actual_EV_Percent": total / n if n else 0.0,
                "Max_Drawdown_Percent": bw.calculate_max_drawdown(tr),
            })
            if total > best_pnl:
                best_pnl, best_log, best_win = total, logs, (train_size, fwd_size)

    return grid_results, best_log, best_win


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def main():
    group = bw.choose_currency_group()
    if not group:
        return
    entry_time = bw.choose_entry_time()
    settings = bw.choose_grid_settings()
    min_train, max_train, min_fwd, max_fwd, _, _ = settings

    parsed = paste_news(group, entry_time)
    if parsed is None:
        return
    news_df, has_surprise = parsed

    pair_arrays_map = {}
    for pair in bw.pair_sets[group]:
        df = bw.load_csv(pair)
        if df is not None:
            pair_arrays_map[pair] = bw.build_pair_arrays(pair, df)
    if not pair_arrays_map:
        print("No valid pair data found.")
        return

    results = {}  # label -> (grid_results, best_log, best_window)
    # allow filtering via prompt (optional)
    mode_set = None
    print("\n--- Mode set ---")
    print("  1) Three modes")
    print("  2) Candle only")
    print("  3) Surprise only")
    raw = input("Choose mode set (default 1): ").strip()
    if raw == "2":
        mode_set = [m for m in MODES if m[0] == "candle_colour"]
    elif raw == "3":
        mode_set = [m for m in MODES if m[0] == "surprise"]
    else:
        mode_set = list(MODES)
    for mode, inverted, label in mode_set:
        if mode == "surprise" and not has_surprise:
            continue
        ev = prepare_events(news_df, mode, inverted, pair_arrays_map)
        print(f"\n=== {label}: {len(ev)} usable events ===")
        if len(ev) == 0:
            continue
        out = run_grid(mode, ev, group, pair_arrays_map, settings)
        if out[0]:
            results[label] = out
            print(f"{label}: best window (train/fwd) = {out[2]}")
        else:
            print(f"{label}: no valid grid results.")

    if not results:
        print("Nothing to save.")
        return

    bw.OUTPUT_DIR.mkdir(exist_ok=True)
    out_path = bw.OUTPUT_DIR / f"{group}_{time.strftime('%Y%m%d_%H%M%S')}_three_modes_{('candle' if mode_set and len(mode_set)==1 and mode_set[0][0]=='candle_colour' else ('surprise' if mode_set and all(m[0]=='surprise' for m in mode_set) else 'three'))}_feeless.xlsx"

    train_idx = list(range(min_train, max_train + 1))
    fwd_cols = list(range(min_fwd, max_fwd + 1))
    rule = ColorScaleRule(
        start_type="num", start_value=-10, start_color="F8696B",
        mid_type="num", mid_value=0, mid_color="FFFFFF",
        end_type="num", end_value=10, end_color="63BE7B",
    )

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # --- all heatmaps on ONE sheet, side by side ---
        sheet = "PnL_Heatmaps"
        start_col = 0
        for label, (grid, _, best_win) in results.items():
            hm = (pd.DataFrame(grid)
                  .pivot(index="Train_Size", columns="Forward_Size", values="Actual_Total_PnL_Percent")
                  .reindex(index=train_idx, columns=fwd_cols))
            hm.to_excel(writer, sheet_name=sheet, startrow=2, startcol=start_col)
            ws = writer.sheets[sheet]
            title = ws.cell(row=1, column=start_col + 1, value=f"{label}  (rows=Train, cols=Forward, best={best_win})")
            title.font = Font(bold=True)
            c1 = openpyxl.utils.get_column_letter(start_col + 2)
            c2 = openpyxl.utils.get_column_letter(start_col + 1 + len(fwd_cols))
            ws.conditional_formatting.add(f"{c1}4:{c2}{3 + len(train_idx)}", rule)
            start_col += len(fwd_cols) + 3  # gap of 2 blank columns

        # --- summaries & logs ---
        for label, (grid, log, _) in results.items():
            (pd.DataFrame(grid).sort_values("Actual_Total_PnL_Percent", ascending=False)
             .to_excel(writer, sheet_name=f"Summary_{label}"[:31], index=False))
        for label, (_, log, _) in results.items():
            pd.DataFrame(log).to_excel(writer, sheet_name=f"Log_{label}"[:31], index=False)

    print("\n--- FINISHED ---")
    print(f"File saved: {out_path}")


if __name__ == "__main__":
    main()
